from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from docx import Document as DocxDocument
from pypdf import PdfReader


SUPPORTED_EXTENSIONS = {
    ".pdf": "PDF",
    ".docx": "DOCX",
    ".txt": "TXT",
    ".md": "MARKDOWN",
    ".markdown": "MARKDOWN",
}


@dataclass
class ExtractedPage:
    """Text extracted from one logical document page."""

    text: str
    page: int | None = None


@dataclass
class ExtractedDocument:
    """Result of document extraction."""

    filename: str
    document_type: str
    pages: list[ExtractedPage]

    @property
    def full_text(self) -> str:
        """Return all extracted text as one string."""
        return "\n\n".join(
            page.text.strip()
            for page in self.pages
            if page.text.strip()
        )


def get_document_type(filename: str) -> str:
    """
    Return the supported document format for a filename.

    Raises:
        ValueError: If the extension is not supported.
    """
    extension = Path(filename).suffix.lower()

    if extension not in SUPPORTED_EXTENSIONS:
        supported = ", ".join(sorted(SUPPORTED_EXTENSIONS))
        raise ValueError(
            f"Unsupported document type '{extension}'. "
            f"Supported extensions are: {supported}"
        )

    return SUPPORTED_EXTENSIONS[extension]


def extract_text(file_path: str | Path, filename: str | None = None) -> ExtractedDocument:
    """
    Extract text from PDF, DOCX, TXT or Markdown.

    PDF page numbers are preserved where available.
    TXT and Markdown are treated as plain text.
    DOCX content is returned as a single logical page.
    """
    path = Path(file_path)

    if not path.exists():
        raise FileNotFoundError(f"Document does not exist: {path}")

    original_filename = filename or path.name
    document_type = get_document_type(original_filename)

    if document_type == "PDF":
        pages = _extract_pdf(path)

    elif document_type == "DOCX":
        pages = _extract_docx(path)

    elif document_type in {"TXT", "MARKDOWN"}:
        pages = _extract_text_file(path)

    else:
        raise ValueError(f"Unsupported document type: {document_type}")

    if not any(page.text.strip() for page in pages):
        raise ValueError("The document does not contain extractable text.")

    return ExtractedDocument(
        filename=original_filename,
        document_type=document_type,
        pages=pages,
    )


def _extract_pdf(path: Path) -> list[ExtractedPage]:
    """Extract text page-by-page from a PDF."""
    reader = PdfReader(str(path))

    pages: list[ExtractedPage] = []

    for page_number, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""

        pages.append(
            ExtractedPage(
                text=text,
                page=page_number,
            )
        )

    return pages


def _extract_docx(path: Path) -> list[ExtractedPage]:
    """Extract paragraphs and tables from a DOCX document."""
    document = DocxDocument(str(path))

    parts: list[str] = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()

        if text:
            parts.append(text)

    for table in document.tables:
        for row in table.rows:
            cells = [
                cell.text.strip()
                for cell in row.cells
            ]

            row_text = " | ".join(
                cell for cell in cells if cell
            )

            if row_text:
                parts.append(row_text)

    return [
        ExtractedPage(
            text="\n".join(parts),
            page=None,
        )
    ]


def _extract_text_file(path: Path) -> list[ExtractedPage]:
    """Read TXT or Markdown as plain UTF-8 text."""
    raw_bytes = path.read_bytes()

    try:
        text = raw_bytes.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = raw_bytes.decode("utf-8", errors="replace")

    return [
        ExtractedPage(
            text=text,
            page=None,
        )
    ]